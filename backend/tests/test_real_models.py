"""Opt-in real weights: RUN_MODEL_TESTS=1; downloads several GB on first run."""
import os
import pytest

pytestmark = pytest.mark.skipif(os.getenv('RUN_MODEL_TESTS') != '1', reason='Opt-in BGE model integration')


def test_real_bge_docling_qdrant_end_to_end(tmp_path):
    from docx import Document
    from app.core.config import Settings
    from app.services.container import Services
    settings = Settings(_env_file=None, document_dir=tmp_path / 'docs',
        data_dir=tmp_path / 'data', qdrant_path=':memory:')
    settings.document_dir.mkdir()
    for name, heading, text in [
        ('supplier.docx', 'Supplier evaluation', 'Suppliers must be evaluated annually by the purchasing manager.'),
        ('calibration.docx', 'Calibration', 'Measurement equipment must be calibrated every six months.')]:
        doc = Document()
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
        doc.save(settings.document_dir / name)
    services = Services(settings)
    try:
        report = services.ingestion.ingest_folder()
        assert all(entry['status'] == 'indexed' for entry in report['files'])
        response = services.search.search('How often must suppliers be evaluated?', top_k=2)
        assert response.results[0].metadata['document_name'] == 'supplier.docx'
        assert 'annually' in response.results[0].text
        assert services.embeddings.dimension == 1024
        model = services.embeddings.model
        reranker = services.reranker._model
        services.search.search('When is equipment calibrated?')
        assert services.embeddings.model is model
        assert services.reranker._model is reranker
    finally:
        services.close()
