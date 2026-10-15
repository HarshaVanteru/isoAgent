"""Real Docling/HybridChunker and embedded Qdrant; no model downloads."""
from io import BytesIO
from types import SimpleNamespace
import shutil
import subprocess

import pytest
from docx import Document
from qdrant_client import QdrantClient
from tokenizers import Tokenizer
from tokenizers.models import WordLevel
from tokenizers.pre_tokenizers import Whitespace
from transformers import PreTrainedTokenizerFast

from app.core.config import Settings
from app.core.errors import ConversionError, InvalidFile, IndexConfigurationError, ParsingError
from app.ingestion.converter import OfficeConverter
from app.ingestion.parser import DoclingParser
from app.ingestion.chunker import StructuredChunker
from app.ingestion.metadata import document_type
from app.models.responses import SearchHit
from app.retrieval.retriever import Retriever
from app.services.ingestion_service import IngestionService
from app.services.search_service import SearchService
from app.vectorstore.qdrant_store import QdrantStore


class TestEmbeddings:
    __test__ = False
    dimension = 3

    def __init__(self):
        vocabulary = ['[UNK]', 'Quality', 'Supplier', 'Control', 'Evaluate', 'suppliers',
                      'annually', 'policy', '.', 'Purchase', 'records', 'Form', 'Name',
                      'Policy', 'Acme', 'Annual']
        tokenizer = Tokenizer(WordLevel({v: i for i, v in enumerate(vocabulary)}, unk_token='[UNK]'))
        tokenizer.pre_tokenizer = Whitespace()
        self.model = SimpleNamespace(tokenizer=PreTrainedTokenizerFast(
            tokenizer_object=tokenizer, unk_token='[UNK]', model_max_length=8192))
        self.calls = 0

    def embed_documents(self, texts):
        self.calls += 1
        return [[float('supplier' in text.lower()), float('quality' in text.lower()), 0.1] for text in texts]

    def embed_query(self, query):
        return self.embed_documents([query])[0]


class TestReranker:
    __test__ = False

    def rerank(self, query, candidates):
        for hit in candidates:
            hit.score = float('supplier' in hit.text.lower())
        return sorted(candidates, key=lambda hit: hit.score, reverse=True)


@pytest.fixture
def pipeline(tmp_path):
    root = tmp_path / 'documents'
    root.mkdir()
    settings = Settings(_env_file=None, document_dir=root, data_dir=tmp_path / 'data',
                        qdrant_path=':memory:', chunk_max_tokens=32)
    embeddings = TestEmbeddings()
    store = QdrantStore(settings, QdrantClient(':memory:'))
    ingestion = IngestionService(settings, OfficeConverter(settings), DoclingParser(),
        StructuredChunker(settings, embeddings), embeddings, store)
    search = SearchService(Retriever(embeddings, store, settings), TestReranker(), settings)
    yield SimpleNamespace(settings=settings, embeddings=embeddings, store=store,
                          ingestion=ingestion, search=search)
    store.close()


def make_doc(path, heading='Supplier Control', text='Evaluate suppliers annually.'):
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(text)
    doc.save(path)
    return path


def test_ingest_search_update_duplicate_and_delete(pipeline):
    p = pipeline
    source = make_doc(p.settings.document_dir / 'XX-OP-253 Supplier Procedure.docx')
    first = p.ingestion.ingest_file(source)
    assert first.status == 'indexed'
    calls = p.embeddings.calls
    assert p.ingestion.ingest_file(source).status == 'duplicate'
    assert p.embeddings.calls == calls
    results = p.search.search('supplier evaluation').results
    assert 'Evaluate suppliers annually' in results[0].text
    assert results[0].metadata['document_code'] == 'XX-OP-253'
    assert results[0].metadata['section'] == 'Supplier Control'
    assert results[0].citation == 1
    assert p.search.search('supplier', filters={'category': 'missing'}).results == []
    make_doc(source, text='Retain supplier records.')
    second = p.ingestion.ingest_file(source)
    assert first.document_id != second.document_id
    assert p.store.document(first.document_id) is None
    assert p.ingestion.delete_document(second.document_id)
    assert p.store.count() == 0


def test_duplicate_alias_and_folder_reconciliation(pipeline):
    p = pipeline
    first = make_doc(p.settings.document_dir / 'a.docx')
    second = p.settings.document_dir / 'b.docx'
    shutil.copyfile(first, second)
    report = p.ingestion.ingest_folder()
    assert len(report['files']) == 2
    assert report['chunks'] == 1
    first.unlink()
    assert p.ingestion.ingest_folder()['chunks'] == 1
    assert p.search.search('supplier').results[0].metadata['source'] == 'b.docx'
    second.unlink()
    assert p.ingestion.ingest_folder()['chunks'] == 0


def test_upload_sanitization_and_limits(pipeline):
    p = pipeline
    response = p.ingestion.upload(BytesIO(b'Evaluate suppliers annually.'), '../../policy.txt')
    assert response.document_name == 'policy.txt'
    assert list((p.settings.data_dir / 'raw').rglob('policy.txt'))
    assert p.ingestion.ingest_folder()['chunks'] == response.chunks
    p.settings.upload_max_bytes = 2
    with pytest.raises(InvalidFile, match='exceeds'):
        p.ingestion.upload(BytesIO(b'large'), 'policy.txt')
    with pytest.raises(InvalidFile):
        p.ingestion.ingest_folder(p.settings.data_dir)


def test_invalid_empty_corrupt_and_unsupported(pipeline):
    p = pipeline
    for name, content in [('empty.txt', b''), ('bad.exe', b'bad'), ('broken.docx', b'bad')]:
        (p.settings.document_dir / name).write_bytes(content)
    report = p.ingestion.ingest_folder()
    assert all(entry['status'] == 'error' for entry in report['files'])
    assert p.store.count() == 0
    assert p.ingestion.report_path.exists()


def test_conversion_separate_and_original_unchanged(pipeline, monkeypatch):
    p = pipeline
    path = p.settings.document_dir / 'old.xls'
    path.write_bytes(b'legacy excel')
    p.settings.libreoffice_path = 'fake-soffice'
    def convert(command, **kwargs):
        assert command[command.index('--convert-to') + 1] == 'xlsx'
        assert 'converted' not in command[1]
        assert '.office-work' in command[1]
        from pathlib import Path
        out = Path(command[command.index('--outdir') + 1])
        (out / 'old.xlsx').write_bytes(b'converted')
        return subprocess.CompletedProcess(command, 0)
    monkeypatch.setattr(subprocess, 'run', convert)
    result = p.ingestion.converter.convert(path)
    assert result.suffix == '.xlsx'
    assert path.read_bytes() == b'legacy excel'
    def fail(*args, **kwargs):
        raise subprocess.TimeoutExpired('soffice', 1)
    monkeypatch.setattr(subprocess, 'run', fail)
    with pytest.raises(ConversionError):
        p.ingestion.converter.convert(path)


def test_ooxml_with_doc_extension(pipeline):
    p = pipeline
    modern = make_doc(p.settings.document_dir / 'modern.docx')
    legacy = modern.with_suffix('.doc')
    shutil.copyfile(modern, legacy)
    assert p.ingestion.ingest_file(legacy).chunks > 0


def test_collection_model_mismatch(pipeline):
    p = pipeline
    p.ingestion.ingest_file(make_doc(p.settings.document_dir / 'a.docx'))
    other = Settings(_env_file=None, document_dir=p.settings.document_dir,
                     data_dir=p.settings.data_dir, qdrant_path=':memory:', embedding_model='different')
    store = QdrantStore(other, p.store.client)
    with pytest.raises(IndexConfigurationError):
        store.ensure_collection(3)
    with pytest.raises(IndexConfigurationError):
        store.ensure_collection(12)


def test_embedded_qdrant_persists_after_reopen(tmp_path):
    from app.models.responses import Chunk
    settings = Settings(_env_file=None, data_dir=tmp_path / 'data',
                        qdrant_path=str(tmp_path / 'qdrant'))
    store = QdrantStore(settings)
    store.ensure_collection(3)
    store.insert([Chunk(text='Supplier policy', metadata={
        'document_id': 'stable', 'chunk_index': 0, 'document_chunk_count': 1})], [[1.0, 0.0, 0.0]])
    store.close()
    reopened = QdrantStore(settings)
    try:
        reopened.ensure_collection(3)
        assert reopened.document('stable')['chunks'] == 1
        assert reopened.search([1.0, 0.0, 0.0], 5, {})[0].text == 'Supplier policy'
    finally:
        reopened.close()


def test_spreadsheet_headers_rows_and_sheet(pipeline):
    from openpyxl import Workbook
    p = pipeline
    path = p.settings.document_dir / 'register.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Suppliers'
    sheet.append(['Supplier', 'Quality'])
    for i in range(30):
        sheet.append([f'Supplier {i}', 'Evaluate annually'])
    book.save(path)
    result = p.ingestion.ingest_file(path)
    assert result.chunks > 1
    hits = p.search.search('supplier').results
    assert all(h.metadata.get('sheet') == 'Suppliers' for h in hits)
    assert all('Supplier' in h.text for h in hits)
    assert all(h.metadata['row_start'] <= h.metadata['row_end'] for h in hits)


def test_spreadsheet_original_row_numbers_and_singletons(pipeline):
    from openpyxl import Workbook
    p = pipeline
    path = p.settings.document_dir / 'offset.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Quality'
    sheet['C10'] = 'Supplier'
    sheet['D10'] = 'Quality'
    sheet['C11'] = 'Acme'
    sheet['D11'] = 'Annual'
    sheet['A1'] = 'Policy'
    book.save(path)
    p.ingestion.ingest_file(path)
    hits = p.search.search('supplier').results
    table_hits = [h for h in hits if 'row_start' in h.metadata]
    assert table_hits[0].metadata['row_start'] == 11
    assert any('Policy' in h.text for h in hits)


def test_model_failure_does_not_delete_prior_document(pipeline, monkeypatch):
    from app.core.errors import ModelError
    p = pipeline
    path = make_doc(p.settings.document_dir / 'policy.docx')
    result = p.ingestion.ingest_file(path)
    make_doc(path, text='New supplier policy')
    def unavailable(texts):
        raise ModelError('Model unavailable')
    monkeypatch.setattr(p.embeddings, 'embed_documents', unavailable)
    with pytest.raises(ModelError):
        p.ingestion.ingest_folder()
    assert p.store.exists(result.document_id)


def test_partial_insert_is_cleaned_and_can_retry(pipeline, monkeypatch):
    from app.core.errors import StorageError
    p = pipeline
    path = make_doc(p.settings.document_dir / 'policy.docx', text='Supplier policy. ' * 60)
    original = p.store.insert
    def partial(chunks, vectors):
        original(chunks[:1], vectors[:1])
        raise StorageError('Insertion failed')
    monkeypatch.setattr(p.store, 'insert', partial)
    with pytest.raises(StorageError):
        p.ingestion.ingest_file(path)
    assert p.store.count() == 0
    monkeypatch.setattr(p.store, 'insert', original)
    assert p.ingestion.ingest_file(path).chunks > 1


def test_model_services_load_once_and_rerank(pipeline, monkeypatch):
    import numpy as np
    import sentence_transformers
    from app.embeddings.bge_embeddings import BGEEmbeddings
    from app.retrieval.reranker import BGEReranker
    created = []
    class Encoder:
        def __init__(self, name, **kwargs):
            created.append(name)
        def encode(self, texts, **kwargs):
            return np.array([[1.0, 0.0, 0.0] for _ in texts])
        def get_embedding_dimension(self):
            return 3
    class CrossEncoder:
        def __init__(self, name, **kwargs):
            created.append(name)
        def predict(self, pairs, **kwargs):
            return np.array([1.0 if 'supplier' in pair[1] else -1.0 for pair in pairs])
    monkeypatch.setattr(sentence_transformers, 'SentenceTransformer', Encoder)
    monkeypatch.setattr(sentence_transformers, 'CrossEncoder', CrossEncoder)
    embeddings = BGEEmbeddings(pipeline.settings)
    assert embeddings.dimension == 3
    embeddings.embed_documents(['a', 'b'])
    assert embeddings.embed_query('c') == [1.0, 0.0, 0.0]
    reranker = BGEReranker(pipeline.settings)
    hits = [SearchHit(text=t, score=0.5, vector_score=0.5) for t in ['calibration', 'supplier']]
    assert reranker.rerank('supplier', hits)[0].text == 'supplier'
    reranker.rerank('supplier', hits)
    assert created == [pipeline.settings.embedding_model, pipeline.settings.reranker_model]


def test_small_form_and_token_budget(pipeline):
    p = pipeline
    form = make_doc(p.settings.document_dir / 'OF-Q-001 Form.docx', 'Form', 'Name Quality')
    assert document_type(form) == 'form'
    assert p.ingestion.ingest_file(form).chunks == 1
    long = make_doc(p.settings.document_dir / 'long.docx', text='Evaluate suppliers annually. ' * 80)
    p.ingestion.ingest_file(long)
    hits = p.search.search('supplier', top_k=100).results
    assert all(len(p.embeddings.model.tokenizer.encode(h.text, add_special_tokens=False)) <= 32 for h in hits)


def test_slides_never_merge_across_boundaries(pipeline):
    from pptx import Presentation
    p = pipeline
    path = p.settings.document_dir / 'training.pptx'
    deck = Presentation()
    for text in ('Supplier', 'Quality'):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = text
    deck.save(path)
    response = p.ingestion.ingest_file(path)
    assert response.chunks == 2
    hits = p.search.search('supplier').results
    assert {h.metadata['slide'] for h in hits} == {1, 2}
    assert not any('Supplier' in h.text and 'Quality' in h.text for h in hits)


def test_api_local_search_and_document_lifecycle(pipeline):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.api.dependencies import get_services
    p = pipeline
    app.dependency_overrides[get_services] = lambda: p
    try:
        with TestClient(app) as client:
            response = client.post('/documents/ingest', files={'file': ('policy.txt', b'Evaluate suppliers annually.')})
            assert response.status_code == 200, response.text
            document_id = response.json()['document_id']
            assert client.get(f'/documents/{document_id}').status_code == 200
            assert client.post('/search', json={'query': 'supplier'}).json()['results']
            assert client.post('/search', json={'query': '  '}).status_code == 422
            assert client.post('/search', json={'query': 'supplier', 'filters': {'bad': 'x'}}).status_code == 422
            assert client.post('/documents/ingest-folder', json={'path': str(p.settings.data_dir)}).status_code == 400
            assert client.delete(f'/documents/{document_id}').status_code == 200
            assert client.get(f'/documents/{document_id}').status_code == 404
            assert client.get('/health').status_code == 200
    finally:
        app.dependency_overrides.clear()


def test_graph_uses_real_qdrant_candidates_and_reranks_union_once(pipeline, monkeypatch):
    from langchain_core.messages import AIMessage
    from langgraph.checkpoint.memory import MemorySaver
    from app.conversation.graph import build_graph
    from app.services.document_index import DocumentIndex
    from tests.test_graph import FakeLLM
    p = pipeline
    p.ingestion.ingest_file(make_doc(p.settings.document_dir / 'supplier.docx'))
    index = object.__new__(DocumentIndex)
    index.settings, index.services = p.settings, p
    calls = []
    rerank = p.search.reranker.rerank
    def counted(query, candidates):
        calls.append(len(candidates))
        return rerank(query, candidates)
    monkeypatch.setattr(p.search.reranker, 'rerank', counted)
    class GroundedLLM(FakeLLM):
        def invoke(self, messages):
            assert 'Evaluate suppliers annually' in messages[0].content
            return AIMessage(content='Evaluate suppliers annually [1].')
    graph = build_graph(index, GroundedLLM(), MemorySaver(), p.settings)
    embedding_calls = p.embeddings.calls
    result = graph.invoke({'question': 'Supplier evaluation?'}, {'configurable': {'thread_id': 'real-store'}})
    assert calls == [1]
    assert p.embeddings.calls == embedding_calls + 1
    assert len(result['sources']) == 1
    assert result['sources'][0]['metadata']['document_name'] == 'supplier.docx'


def test_rank_fusion_deduplicates_across_queries(pipeline):
    from app.services.search_service import SearchService
    def hit(identity, score):
        return SearchHit(text=identity, score=score, vector_score=score, metadata={'chunk_id': identity})
    class Retriever:
        def retrieve(self, query, limit, filters):
            assert filters == {'category': 'Quality'}
            return [hit('a', .9), hit('shared', .7)] if query == 'first' else [hit('shared', .8), hit('b', .9)]
    service = SearchService(Retriever(), TestReranker(), pipeline.settings)
    candidates = service.retrieve_candidates(['first', 'second'], {'category': 'Quality'})
    assert len(candidates) == 3
    assert candidates[0].metadata['chunk_id'] == 'shared'
    assert candidates[0].vector_score == .8
